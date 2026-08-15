import os
import sys
# pyrefly: ignore [missing-import]
from pptx import Presentation

# Force stdout to use utf-8
sys.stdout.reconfigure(encoding='utf-8')

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
PPTX_PATH = os.path.join(PROJECT_ROOT, "PredictaGuard_Presentation.pptx")
OUTPUT_PATH = os.path.join(PROJECT_ROOT, "ml", "scratch", "presentation_text.txt")

def read_presentation():
    if not os.path.exists(PPTX_PATH):
        print(f"Presentation file not found at: {PPTX_PATH}")
        return
        
    prs = Presentation(PPTX_PATH)
    print(f"Loaded presentation: {PPTX_PATH}")
    print(f"Total slides: {len(prs.slides)}")
    
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        f.write(f"Loaded presentation: {PPTX_PATH}\n")
        f.write(f"Total slides: {len(prs.slides)}\n\n")
        
        for i, slide in enumerate(prs.slides):
            f.write(f"--- Slide {i+1} ---\n")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            text = run.text.strip()
                            if text:
                                f.write(f"  Shape: '{text}'\n")
                if shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            text = cell.text.strip()
                            if text:
                                f.write(f"  Table: '{text}'\n")
            f.write("\n")
            
    print(f"Dumped to {OUTPUT_PATH}")

if __name__ == "__main__":
    read_presentation()
